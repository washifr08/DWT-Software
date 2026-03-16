from customtkinter import * 
from tkinter import messagebox  
from openpyxl import *  
import os  
import sys  
import subprocess  
from matplotlib.figure import Figure 
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg 
from datetime import datetime, timedelta 

current_page = "macro_calc"
update_chart_progress_theme = None

macro_calc = CTk()
macro_calc.geometry("1200x800")
macro_calc.title("D.W.T - Macro Nutritional Calculator")

set_default_color_theme("green")

screen_width = macro_calc.winfo_screenwidth()
screen_height = macro_calc.winfo_screenheight()

centre_x = int((screen_width - 1200) / 2)
centre_y = int((screen_height - 1000) / 2)

macro_calc.geometry(f"+{centre_x}+{centre_y}")

header = CTkFrame(macro_calc, fg_color=("gray90", "#212121"), height=50, corner_radius=0)
header.pack(fill="both", side="top")

header_title = CTkLabel(header, text="Digital Wellbeing Tracker", font=("Century Gothic", 22, "bold"))
header_title.pack(side="left", padx=20)

def dark_light_toggle():
    theme_state = get_appearance_mode()
    if theme_state == "Dark":
        set_appearance_mode("light")
        set_all_light()
    else:
        set_appearance_mode("dark")
        set_all_dark()
    
    global current_page, update_chart_progress_theme

    if current_page == "macro_calc":
        update_chart_theme()
    elif current_page == "progress":
        if update_chart_progress_theme is not None:
            update_chart_progress_theme()

theme_switch = CTkSwitch(header, text="Dark / Light Mode", font=("Century Gothic", 14), command=dark_light_toggle)
theme_switch.pack(side="right", padx=20, pady=10)

contents_frame = CTkFrame(macro_calc, fg_color=("gray95", "#212121"), height=200, corner_radius=25)
contents_frame.pack(fill="both", expand=True, pady=(25, 20), padx=30)

def set_all_light():
    macro_calc.configure(fg_color="gray85")
    header.configure(fg_color="gray95")
    contents_frame.configure(fg_color="gray95")

def set_all_dark():
    macro_calc.configure(fg_color="#1A1A1A")
    header.configure(fg_color="#212121")
    contents_frame.configure(fg_color="#212121")

def update_chart_theme():
    global chart_drawn
    if get_appearance_mode() == "Dark":
        fig.patch.set_facecolor("#212121")
        ax.set_facecolor("#212121")
        title_color = "white"
        text_color = "white"
    else:
        fig.patch.set_facecolor("#f2f2f2")
        ax.set_facecolor("#f2f2f2")
        title_color = "black"
        text_color = "black"

    if chart_drawn:
        ax.set_title("Macros (g)", color=title_color, fontfamily="Century Gothic", fontsize=14)
        for text in ax.texts:
            text.set_color(text_color)
    else:
        ax.set_title("")

    canvas.draw()

if get_appearance_mode() == "Dark":
    set_all_dark()
else:
    set_all_light()

if len(sys.argv) >= 3:
    user_email = sys.argv[2]
    user_name = sys.argv[1]
else:
    user_email = "test"

def make_safe_filename(email):
    email = email.strip().lower()
    email = email.replace("@", "_at_")
    email = email.replace(".", "_dot_")
    return email

def get_user_workbook_path():
    basedir = os.path.dirname(os.path.abspath(__file__))
    folder = os.path.join(basedir, "user_databases")
    if not os.path.exists(folder):
        os.makedirs(folder)

    safe_email = make_safe_filename(user_email)
    return os.path.join(folder, f"{safe_email}.xlsx")

def create_user_workbook_if_new():
    
    path = get_user_workbook_path()
    
    if not os.path.exists(path):
        wb = Workbook()
        ws = wb.active
        ws.title = "Macro Progress"
        ws.append(["Date", "Weight (kg)", "Calories (kcal)", "Protein (g)", "Fats (g)", "Carbohydrates (g)"])
        wb.save(path)
        wb.close()

    return path

def save_macro_progress(weight, calories, protein, fats, carbs):

    path = create_user_workbook_if_new()
    wb = load_workbook(path)
    ws = wb["Macro Progress"]

    from datetime import datetime
    today = datetime.now().strftime("%Y-%m-%d")

    ws.append([today, round(weight, 1), int(calories), int(protein), int(fats), int(carbs)])
    wb.save(path)
    wb.close()

def open_progress_window():

    path = create_user_workbook_if_new()
    wb = load_workbook(path)
    ws = wb["Macro Progress"]

    week_weights = {}

    for row in ws.iter_rows(min_row=2, values_only=True):
        date_str = row[0]
        weight = row[1]
        if date_str != "" and weight != "":
            date = datetime.strptime(date_str, "%Y-%m-%d")

            # find the start of the week (monday)
            week_start = date - timedelta(days=date.weekday())

            if week_start not in week_weights:
                week_weights[week_start] = []

            week_weights[week_start].append(float(weight))

    if len(week_weights) == 0:
        messagebox.showinfo("No Data", "No weight data available to show progress.")
        return
    
    # sort weeks and calculate mean weifghts
    sorted_weeks = sorted(week_weights.keys())
    week_labels = []
    mean_weights = []

    for w in sorted_weeks:
        weights_lists = week_weights[w]
        mean_weight = sum(weights_lists) / len(weights_lists)
        week_labels.append(w.strftime("%d %b"))
        mean_weights.append(mean_weight)

    global contents_frame, current_page, update_chart_progress_theme

    contents_frame.destroy()

    contents_frame = CTkFrame(macro_calc, fg_color=("gray95", "#212121"), height=200, corner_radius=25)
    contents_frame.pack(fill="both", expand=True, pady=(35, 20), padx=30)

    current_page = "progress"

    if get_appearance_mode() == "Dark":
        set_all_dark()
    else:
        set_all_light()
    
    progress_title = CTkLabel(contents_frame, text="Weekly Weight Progress",
                           font=("Century Gothic", 30, "bold"), text_color="#4CAF50")
    progress_title.pack(pady=(30, 0))

    line_graph_frame = CTkFrame(contents_frame, fg_color=("gray95", "#212121"), corner_radius=20, width=950, height=560)
    line_graph_frame.pack(pady=0)

    fig2 = Figure(figsize=(9.2, 5.0), dpi=100)
    ax2 = fig2.add_subplot(111)

    ax2.plot(week_labels, mean_weights, marker='o')
    ax2.set_xlabel("Weeks (Mon-Sun)", fontfamily="Century Gothic", fontsize=14)
    ax2.set_ylabel("Weight (kg)", fontfamily="Century Gothic", fontsize=14)
    ax2.set_title("", fontfamily="Century Gothic", fontsize=16)

    for label in ax2.get_xticklabels():
        label.set_rotation(30)

    fig2.subplots_adjust(bottom=0.22)  

    canvas2 = FigureCanvasTkAgg(fig2, master=line_graph_frame)
    canvas_widget2 = canvas2.get_tk_widget()
    canvas_widget2.pack(fill="both", expand=True, padx=15, pady=15)

    def update_chart_progress_theme():

        if get_appearance_mode() == "Dark":
            fig2.patch.set_facecolor("#212121")
            ax2.set_facecolor("#212121")
            ax2.title.set_color("white")
            ax2.xaxis.label.set_color("white")
            ax2.yaxis.label.set_color("white")
            ax2.tick_params(axis='x', colors='white')
            ax2.tick_params(axis='y', colors='white')

            for spine in ax2.spines.values():
                spine.set_color("#f2f2f2")
            
        else:
            fig2.patch.set_facecolor("#f2f2f2")
            ax2.set_facecolor("#f2f2f2")
            ax2.title.set_color("black")
            ax2.xaxis.label.set_color("black")
            ax2.yaxis.label.set_color("black")
            ax2.tick_params(axis='x', colors='black')
            ax2.tick_params(axis='y', colors='black')

            for spine in ax2.spines.values():
                spine.set_color("#212121")
        
        canvas2.draw()
    
    update_chart_progress_theme()
    update_chart_progress_theme = update_chart_progress_theme

    go_back = CTkButton(contents_frame, text="Go Back to Homepage",
                        corner_radius=25, font=("Century Gothic", 16),
                        fg_color="#4CAF50", hover_color="#3E803F",
                        border_color="#000000", border_width=2,
                        text_color="#000000", command=go_back_to_homepage)
    go_back.pack(pady=(10, 20))

def go_back_to_homepage():
    macro_calc.after(100, macro_calc.destroy)
    subprocess.Popen([sys.executable, "homepage.py", user_name, user_email]) 

title_label = CTkLabel(contents_frame, text="Your personalised Macro Nutritional Calculator",
                       font=("Century Gothic", 30, "bold"), text_color="#4CAF50")
title_label.place(relx=0.5, y=18, anchor="n")

weight_entry_text = CTkLabel(contents_frame, text="Weight (kg)", font=("Century Gothic", 14))
weight_entry_text.pack(anchor="w", padx=25, pady=(70, 2))
weight_entry = CTkEntry(contents_frame, placeholder_text="e.g. 62.5", width=300, height=36,
                        corner_radius=10, font=("Century Gothic", 16))
weight_entry.pack(anchor="w", padx=25, pady=(0, 1))

height_entry_text = CTkLabel(contents_frame, text="Height (cm)", font=("Century Gothic", 14))
height_entry_text.pack(anchor="w", padx=25, pady=(16, 2))
height_entry = CTkEntry(contents_frame, placeholder_text="e.g. 175", width=300, height=36,
                        corner_radius=10, font=("Century Gothic", 16))
height_entry.pack(anchor="w", padx=25, pady=(0, 1))

age_entry_text = CTkLabel(contents_frame, text="Age", font=("Century Gothic", 14))
age_entry_text.pack(anchor="w", padx=25, pady=(16, 2))
age_entry = CTkEntry(contents_frame, placeholder_text="e.g. 25", width=300, height=36,
                     corner_radius=10, font=("Century Gothic", 16))
age_entry.pack(anchor="w", padx=25, pady=(0, 1))

sex_label = CTkLabel(contents_frame, text="Sex", font=("Century Gothic", 14))
sex_label.pack(anchor="w", padx=25, pady=(10, 2))
user_sex = StringVar(value="Male")
sex_dropdown = CTkOptionMenu(contents_frame, values=["Male", "Female"], variable=user_sex,
                             width=300, height=36, corner_radius=10, font=("Century Gothic", 16), fg_color="#469E4A",
                             button_color="#3D8B40", button_hover_color="#2F5F30", text_color="#000000")
sex_dropdown.pack(anchor="w", padx=25, pady=(0, 5))

algotihm_label = CTkLabel(contents_frame, text="Choose your algorithm", font=("Century Gothic", 14))
algotihm_label.pack(anchor="w", padx=25, pady=(10, 2))
user_algorithm = StringVar(value="Mifflin-St Jeor")
algorithm_dropdown = CTkOptionMenu(contents_frame, values=["Mifflin-St Jeor", "Katch-McArdle"], variable=user_algorithm,
                                   width=300, height=36, corner_radius=10, font=("Century Gothic", 16), fg_color="#469E4A",
                                   button_color="#3D8B40", button_hover_color="#2F5F30", text_color="#000000")
algorithm_dropdown.pack(anchor="w", padx=25, pady=(0, 5))

weight_change_label = CTkLabel(contents_frame, text="Choose your weight gain/loss rate PER WEEK (kg)", font=("Century Gothic", 14))
weight_change_label.pack(anchor="w", padx=25, pady=(10, 2))
weight_change_entry = CTkEntry(contents_frame, placeholder_text="e.g. 0.5 or -0.5", width=300, height=36,
                               corner_radius=10, font=("Century Gothic", 16))
weight_change_entry.pack(anchor="w", padx=25, pady=(0, 5))

activity_label_title = CTkLabel(contents_frame, text="Activity level", font=("Century Gothic", 14))
activity_label_title.pack(anchor="w", padx=25, pady=(10, 2))
activity_value_label = CTkLabel(contents_frame, text="Moderate", font=("Century Gothic", 16, "bold"))
activity_value_label.pack(anchor="w", padx=25, pady=(0, 4))

different_activitylevel_names = ["Minimal", "Light", "Moderate", "Very Active"]

def activity_slider(number_of_steps):
    rounded_position = int(round(float(number_of_steps)))
    activity_slider_bar.set(rounded_position)
    activity_value_label.configure(text=different_activitylevel_names[rounded_position])

activity_slider_bar = CTkSlider(contents_frame, from_=0, to=3, number_of_steps=3, command=activity_slider, width=300)
activity_slider_bar.set(2) 
activity_slider_bar.pack(anchor="w", padx=25, pady=(0, 5))

chart_frame = CTkFrame(contents_frame, fg_color=("gray95", "#212121"), corner_radius=20, width=520, height=320)
chart_frame.place(x=480, y=95)  

macro_text_frame = CTkFrame(contents_frame, fg_color=("gray95", "#212121"), corner_radius=20, width=520, height=220)
macro_text_frame.place(x=635, y=430)

macro_text_title = CTkLabel(macro_text_frame, text="Macro Results", font=("Century Gothic", 20, "bold"), text_color="#4CAF50")
macro_text_title.pack(pady=(15, 5))

macro_results_label = CTkLabel(macro_text_frame,text="Fill in your details and click 'Calculate your macros!'.",
    font=("Century Gothic", 16), justify="center")
macro_results_label.pack(padx=30, pady=(5, 10))

fig = Figure(figsize=(5.2, 3.2), dpi=100)
ax = fig.add_subplot(111)

chart_drawn = False

fig.patch.set_facecolor("#212121")     
ax.set_facecolor("#212121")            

canvas = FigureCanvasTkAgg(fig, master=chart_frame)
canvas_widget = canvas.get_tk_widget()
canvas_widget.pack(fill="both", expand=True, padx=10, pady=10)

ax.clear()
ax.set_facecolor("#1F1F1F")
ax.axis("off")
canvas.draw()

def is_valid_number(value):
    if value=="" :
        return False
    try:
        float(value)
        return True
    except:
        return False

def calculate_macros():
    weight_string = (weight_entry.get())
    height_string = (height_entry.get())
    age_string = (age_entry.get())
    weight_change_string = (weight_change_entry.get())

    if weight_string == "" or height_string == "" or age_string == "" or weight_change_string == "":
        messagebox.showerror("Input Error", "Please fill in all fields.")
        return
    
    if not is_valid_number(weight_string):
        messagebox.showerror("Input Error", "Please enter a valid number for weight.")
        return

    if not is_valid_number(height_string):
        messagebox.showerror("Input Error", "Please enter a valid number for height.")
        return

    if not is_valid_number(age_string):
        messagebox.showerror("Input Error", "Please enter a valid number for age.")
        return

    if not is_valid_number(weight_change_string):
        messagebox.showerror("Input Error", "Please enter a valid number for weight change.")
        return
    
    weight = float(weight_string)
    height = float(height_string)
    age = float(age_string)
    weight_change = float(weight_change_string)

    # Mifflin-St Jeor calculations
    if user_algorithm.get() == "Mifflin-St Jeor":
        if user_sex.get() == "Male":
            bmr = (10 * weight) + (6.25 * height) - (5 * age) + 5
            if activity_value_label.cget("text") == "Minimal":
                tdee = bmr * 1.2
            elif activity_value_label.cget("text") == "Light":
                tdee = bmr * 1.375
            elif activity_value_label.cget("text") == "Moderate":
                tdee = bmr * 1.55
            else:
                tdee = bmr * 1.725
        else:
            bmr = (10 * weight) + (6.25 * height) - (5 * age) - 161
            if activity_value_label.cget("text") == "Minimal":
                tdee = bmr * 1.2
            elif activity_value_label.cget("text") == "Light":
                tdee = bmr * 1.375
            elif activity_value_label.cget("text") == "Moderate":
                tdee = bmr * 1.55
            else:
                tdee = bmr * 1.725

    # Katch-McArdle calculations
    else:  
        if user_sex.get() == "Male":
            lbm = (0.407 * weight) + (0.267 * height) - 19.2
            bmr = 370 + (21.6 * lbm)
            if activity_value_label.cget("text") == "Minimal":
                tdee = bmr * 1.2
            elif activity_value_label.cget("text") == "Light":
                tdee = bmr * 1.375
            elif activity_value_label.cget("text") == "Moderate":
                tdee = bmr * 1.55
            else:
                tdee = bmr * 1.725
        else:
            lbm = (0.252 * weight) + (0.473 * height) - 48.3
            bmr = 370 + (21.6 * lbm)
            if activity_value_label.cget("text") == "Minimal":
                tdee = bmr * 1.2
            elif activity_value_label.cget("text") == "Light":
                tdee = bmr * 1.375
            elif activity_value_label.cget("text") == "Moderate":
                tdee = bmr * 1.55
            else:
                tdee = bmr * 1.725

    # calculate daily caloric needs including weight change
    # 7700 calories per kg of fat
    surplus_deficit = (weight_change * 7700) / 7  
    daily_caloric_needs = tdee + surplus_deficit

    # 2 grams of protein per kg of body weight
    protein_intake = weight * 2  
    # 28% of calories from fat, 9 calories per gram
    fat_intake = (0.28 * daily_caloric_needs) / 9  
    carb_intake = daily_caloric_needs - ((protein_intake * 4) + (fat_intake * 9))
    # 4 calories per gram of carbohydrate
    carb_intake = carb_intake / 4  

    macro_results_label.configure(text=
        f"Daily calories: {int(daily_caloric_needs)} kcal\n\n"
        f"Protein: {int(protein_intake)} g\n"
        f"Fats: {int(fat_intake)} g\n"
        f"Carbohydrates: {int(carb_intake)} g")

    values = [protein_intake, fat_intake, carb_intake]
    labels = ["Protein", "Fats", "Carbs"]

    save_macro_progress(weight, daily_caloric_needs, protein_intake, fat_intake, carb_intake)

    # safety if something weird happens
    values = [max(0, float(v)) for v in values]
    if sum(values) == 0:
        ax.clear()
        ax.set_title("Macros (grams)")
        ax.text(0.5, 0.5, "No data to display", ha="center", va="center")
        canvas.draw()
        return

    ax.clear()
    ax.set_facecolor("#212121")  

    ax.pie(values, labels=labels, autopct="%1.1f%%", startangle=90,
        textprops={"color": "white", "fontfamily": "Century Gothic", "fontsize": 14})  

    ax.set_title("Macros (g)", color="white", fontfamily="Century Gothic", fontsize=14) 
    ax.axis("equal")
    canvas.draw()

    global chart_drawn
    chart_drawn = True

    update_chart_theme()

def calculate_button_pressed():
    
    calculate_macros()
    view_progress_btn.pack(pady=(20, 0))

calculate_btn = CTkButton(contents_frame, text="Calculate your macros!",
                        corner_radius=25, font=("Century Gothic", 16),
                        fg_color="#4CAF50", hover_color="#3E803F",
                        border_color="#000000", border_width=2,
                        text_color="#000000", command=calculate_button_pressed)
calculate_btn.pack(anchor="w", padx=65, pady=(20, 6))

view_progress_btn = CTkButton(macro_text_frame, text="View Past Progress",
                        corner_radius=25, font=("Century Gothic", 16),
                        fg_color="#4CAF50", hover_color="#3E803F",
                        border_color="#000000", border_width=2,
                        text_color="#000000", command=open_progress_window)

view_progress_btn.pack_forget()

macro_calc.mainloop()
