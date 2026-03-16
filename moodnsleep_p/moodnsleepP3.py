from customtkinter import *
from tkinter import messagebox
from datetime import datetime, timedelta
import sys
import subprocess
import os 
from openpyxl import Workbook, load_workbook

moodnsleep = CTk()
moodnsleep.geometry("1200x800")
moodnsleep.title("D.W.T - Mood and Sleep Tracker")

set_default_color_theme("green")

screen_width = moodnsleep.winfo_screenwidth()
screen_height = moodnsleep.winfo_screenheight()

centre_x = int((screen_width - 1200) / 2)
centre_y = int((screen_height - 1000) / 2)

moodnsleep.geometry(f"+{centre_x}+{centre_y}")

if len(sys.argv) >= 2:
    user_email = sys.argv[1]
else:
    user_email = "test"

header = CTkFrame(moodnsleep, fg_color=("gray90", "#212121"), height=50, corner_radius=0)
header.pack(fill="both", side="top")

header_title = CTkLabel(header, text="Digital Wellbeing Tracker", font=("Century Gothic", 22, "bold"))
header_title.pack(side="left", padx=20)

contents_frame = CTkFrame(moodnsleep, fg_color=("gray95", "#212121"), height=200, corner_radius=25)
contents_frame.pack(fill="both", expand=True, pady=(25, 20), padx=30)

def set_all_light():
    moodnsleep.configure(fg_color="gray85")
    header.configure(fg_color="gray95")
    contents_frame.configure(fg_color="gray95")

def set_all_dark():
    moodnsleep.configure(fg_color="#1A1A1A")
    header.configure(fg_color="#212121")
    contents_frame.configure(fg_color="#212121")

def dark_light_toggle():
    theme_state = get_appearance_mode()
    if theme_state == "Dark":
        set_appearance_mode("light")
        set_all_light()
    else:
        set_appearance_mode("dark")
        set_all_dark()

theme_switch = CTkSwitch(header, text="Dark / Light Mode", font=("Century Gothic", 14), command=dark_light_toggle)
theme_switch.pack(side="right", padx=20, pady=10)

if get_appearance_mode() == "Dark":
    set_all_dark()
else:
    set_all_light()

date_today = datetime.now().strftime("%Y-%m-%d")

def time_to_minutes(hhmm):
    h, m = hhmm.split(":")
    return int(h) * 60 + int(m)

def calc_sleep_hours(bed_time, wake_time):
    bed_mins = time_to_minutes(bed_time)
    wake_mins = time_to_minutes(wake_time)
    if wake_mins <= bed_mins:
        wake_mins += 24 * 60
    mins = wake_mins - bed_mins
    return round(mins / 60, 2)

def get_selected_times():
    bed_time = f"{bed_hour_var.get()}:{bed_minute_var.get()}"
    wake_time = f"{wake_hour_var.get()}:{wake_minute_var.get()}"
    return bed_time, wake_time

def get_today_values():
    mood = int(mood_slider.get())
    bed_time, wake_time = get_selected_times()
    sleep_hours = calc_sleep_hours(bed_time, wake_time)
    sleep_quality = int(sleep_quality_var.get())
    notes = notes_box.get("1.0", "end").strip()
    return mood, bed_time, wake_time, sleep_hours, sleep_quality, notes

def week_start_monday(date_str):
    d = datetime.strptime(date_str, "%Y-%m-%d").date()
    monday = d - timedelta(days=d.weekday())
    return monday

def week_dates_from_monday(monday_date):
    return [(monday_date + timedelta(days=i)).isoformat() for i in range(7)]

def shift_week(monday_date, weeks):
    return monday_date + timedelta(days=7 * weeks)

def mean_ignore_none(values):
    nums = [v for v in values if v is not None]
    if len(nums) == 0:
        return None
    return round(sum(nums) / len(nums), 2)

def build_week_series(week_dates, data_by_date):
    day_names = ("Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun")

    mood_vals = []
    sleep_hours_vals = [] 
    quality_vals = []

    for ds in week_dates:
        if ds in data_by_date:
            mood_vals.append(data_by_date[ds].get("mood"))
            sleep_hours_vals.append(data_by_date[ds].get("sleep_hours"))
            quality_vals.append(data_by_date[ds].get("sleep_quality"))
        else:
            mood_vals.append(None)
            sleep_hours_vals.append(None)
            quality_vals.append(None)

        summary = { "avg_mood": mean_ignore_none(mood_vals),
            "avg_sleep_hours": mean_ignore_none(sleep_hours_vals),
            "avg_sleep_quality": mean_ignore_none(quality_vals) }
    
    return day_names, mood_vals, sleep_hours_vals, quality_vals, summary

current_monday = week_start_monday(date_today)
week_store = {}

title_label = CTkLabel(contents_frame, text="Mood & Sleep Tracker",
                       font=("Century Gothic", 30, "bold"), text_color="#4CAF50")
title_label.place(relx=0.5, y=18, anchor="n")

date_label = CTkLabel(contents_frame, text=f"Entry Date: {date_today}",
                      font=("Century Gothic", 16))
date_label.pack(anchor="w", padx=25, pady=(70, 4))

mood_entry_text = CTkLabel(contents_frame, text="Mood (1–10)", font=("Century Gothic", 14))
mood_entry_text.pack(anchor="w", padx=25, pady=(16, 2))

mood_value_label = CTkLabel(contents_frame, text="5", font=("Century Gothic", 16, "bold"))
mood_value_label.pack(anchor="w", padx=25, pady=(0, 4))

def mood_slider_changed(v):
    mood_value_label.configure(text=str(int(v)))

mood_slider = CTkSlider(contents_frame, from_=1, to=10, number_of_steps=9, command=mood_slider_changed, width=300)
mood_slider.set(5)
mood_slider.pack(anchor="w", padx=25, pady=(0, 5))

sleep_time_title = CTkLabel(contents_frame, text="Sleep time", font=("Century Gothic", 14))
sleep_time_title.pack(anchor="w", padx=25, pady=(14, 2))

times_h = [f"{i:02d}" for i in range(24)]
times_m = [f"{i:02d}" for i in range(0, 60, 5)]

bed_label = CTkLabel(contents_frame, text="Bed time", font=("Century Gothic", 14))
bed_label.pack(anchor="w", padx=25, pady=(6, 2))

bed_row = CTkFrame(contents_frame, fg_color="transparent")
bed_row.pack(anchor="w", padx=25, pady=(0, 6))

bed_hour_var = StringVar(value="23")
bed_minute_var = StringVar(value="00")

bed_hour_menu = CTkOptionMenu(bed_row, values=times_h, variable=bed_hour_var,
                           width=140, height=36, corner_radius=10,
                           font=("Century Gothic", 16), fg_color="#469E4A",
                           button_color="#3D8B40", button_hover_color="#2F5F30",
                           text_color="#000000")
bed_hour_menu.pack(side="left", padx=(0, 10))

bed_minute_menu = CTkOptionMenu(bed_row, values=times_m, variable=bed_minute_var,
                           width=140, height=36, corner_radius=10,
                           font=("Century Gothic", 16), fg_color="#469E4A",
                           button_color="#3D8B40", button_hover_color="#2F5F30",
                           text_color="#000000")
bed_minute_menu.pack(side="left")

wake_label = CTkLabel(contents_frame, text="Wake time", font=("Century Gothic", 14))
wake_label.pack(anchor="w", padx=25, pady=(10, 2))

wake_row = CTkFrame(contents_frame, fg_color="transparent")
wake_row.pack(anchor="w", padx=25, pady=(0, 10))

wake_hour_var = StringVar(value="07")
wake_minute_var = StringVar(value="00")

wake_hour_menu = CTkOptionMenu(wake_row, values=times_h, variable=wake_hour_var,
                            width=140, height=36, corner_radius=10,
                            font=("Century Gothic", 16), fg_color="#469E4A",
                            button_color="#3D8B40", button_hover_color="#2F5F30",
                            text_color="#000000")
wake_hour_menu.pack(side="left", padx=(0, 10))

wake_minute_menu = CTkOptionMenu(wake_row, values=times_m, variable=wake_minute_var,
                            width=140, height=36, corner_radius=10,
                            font=("Century Gothic", 16), fg_color="#469E4A",
                            button_color="#3D8B40", button_hover_color="#2F5F30",
                            text_color="#000000")
wake_minute_menu.pack(side="left")

sleep_quality_label = CTkLabel(contents_frame, text="Sleep quality (1–5)", font=("Century Gothic", 14))
sleep_quality_label.pack(anchor="w", padx=25, pady=(6, 2))

sleep_quality_var = StringVar(value="3")
sleep_quality_menu = CTkOptionMenu(contents_frame, values=["1", "2", "3", "4", "5"], variable=sleep_quality_var,
                                   width=300, height=36, corner_radius=10,
                                   font=("Century Gothic", 16), fg_color="#469E4A",
                                   button_color="#3D8B40", button_hover_color="#2F5F30",
                                   text_color="#000000")
sleep_quality_menu.pack(anchor="w", padx=25, pady=(0, 5))

notes_label = CTkLabel(contents_frame, text="Notes (optional)", font=("Century Gothic", 14))
notes_label.pack(anchor="w", padx=25, pady=(12, 2))

notes_box = CTkTextbox(contents_frame, height=95, width=300)
notes_box.pack(anchor="w", padx=25, pady=(0, 10))

moodsleep_sheet = "MoodSleep"
moodsleep_headers = ["Date", "Mood", "Bed Time", "Wake Time", "Sleep Hours", "Sleep Quality", "Notes"]

def safe_email_filename(email):
    e = email.strip().lower()
    e = e.replace("@", "_at_")
    e = e.replace(".", "_dot_")
    return e

def get_user_workbook_path():
    basedir = os.path.dirname(os.path.abspath(__file__))
    folder = os.path.join(basedir, "user_databases")
    if not os.path.exists(folder):
        os.makedirs(folder)

    safe_name = safe_email_filename(user_email)
    return os.path.join(folder, f"{safe_name}.xlsx")

def ensure_workbook_exists():
    path = get_user_workbook_path()

    if not os.path.exists(path):
        wb = Workbook()
        ws = wb.active
        ws.title = "Macro Progress"
        ws.append(["Date", "Weight (Kg)", "Calories (kcal)", "Protein (g)", "Fats (g)", "Carbohydrates (g)"])
        wb.save(path)
        wb.close()

    return path

def ensure_moodsleep_sheet(path):
    wb = load_workbook(path)
    if moodsleep_sheet not in wb.sheetnames:
        ws = wb.create_sheet(title=moodsleep_sheet)
        ws.append(moodsleep_headers)
        wb.save(path)
    wb.close()

def find_row_by_date(ws, date_str):
    for r in range(2, ws.max_row + 1):
        if ws.cell(row=r, column=1).value == date_str:
            return r
    return None

def save_moodsleep_to_excel(date_str, mood, bed_time, wake_time, sleep_hours, sleep_quality, notes):
    path = ensure_workbook_exists()
    ensure_moodsleep_sheet(path)

    wb = load_workbook(path)
    ws = wb[moodsleep_sheet]

    row = find_row_by_date(ws, date_str)
    if row is None:
        row = ws.max_row + 1

    ws.cell(row=row, column=1).value = date_str
    ws.cell(row=row, column=2).value = int(mood)
    ws.cell(row=row, column=3).value = bed_time
    ws.cell(row=row, column=4).value = wake_time
    ws.cell(row=row, column=5).value = float(sleep_hours)
    ws.cell(row=row, column=6).value = int(sleep_quality)
    ws.cell(row=row, column=7).value = notes

    wb.save(path)
    wb.close()

def load_moodsleep_week_dict(monday_date):
    path = ensure_workbook_exists()
    ensure_moodsleep_sheet(path)

    week_dates = week_dates_from_monday(monday_date)
    wanted = set(week_dates)

    wb = load_workbook(path)
    ws = wb[moodsleep_sheet]

    data = {}
    for r in range(2, ws.max_row + 1):
        d = ws.cell(row=r, column=1).value
        if d in wanted:
            data[d] = { "mood": ws.cell(row=r, column=2).value,
                "sleep_hours": ws.cell(row=r, column=5).value,
                "sleep_quality": ws.cell(row=r, column=6).value, }
            
    wb.close()
    return data

def save_ui():
    mood, bed_time, wake_time, sleep_hours, sleep_quality, notes = get_today_values()
    save_moodsleep_to_excel(date_today, mood, bed_time, wake_time, sleep_hours, sleep_quality, notes)
    
    messagebox.showinfo("Saved", f"Today's entry saved. \nSleep hours: {sleep_hours}")

def view_ui():
    week_data_dict = load_moodsleep_week_dict(current_monday)
    week_dates = week_dates_from_monday(current_monday)

    day_names, mood_vals, sleep_hours_vals, quality_vals, summary = build_week_series(week_dates, week_data_dict)

    msg = ( f"{week_range_text(current_monday)}\n\n"
            f"Mood: {mood_vals}\n"
            f"Sleep hours: {sleep_hours_vals}\n"
            f"Sleep quality: {quality_vals}\n\n"
            f"Averages:\n"
            f"• Mood: {summary['avg_mood']}\n"
            f"• Sleep hours: {summary['avg_sleep_hours']}\n"
            f"• Sleep quality: {summary['avg_sleep_quality']}" )
    
    messagebox.showinfo("Week Data", msg)

save_btn = CTkButton(contents_frame, text="Save Today's Entry",
                     corner_radius=25, font=("Century Gothic", 16),
                     fg_color="#4CAF50", hover_color="#3E803F",
                     border_color="#000000", border_width=2,
                     text_color="#000000", command=save_ui, width=300, height=44)
save_btn.pack(anchor="w", padx=25, pady=(12, 10))

view_btn = CTkButton(contents_frame, text="View Weekly Insights",
                     corner_radius=25, font=("Century Gothic", 16),
                     fg_color="#4CAF50", hover_color="#3E803F",
                     border_color="#000000", border_width=2,
                     text_color="#000000", command=view_ui, width=260, height=44)
view_btn.place(relx=0.7, rely=0.87, anchor="center")

right_top_frame = CTkFrame(contents_frame, fg_color=("gray95", "#212121"), corner_radius=20, width=520, height=320)
right_top_frame.place(relx=0.7, y=120, anchor="n")

right_top_title = CTkLabel(right_top_frame, text="Today’s Overview",
                           font=("Century Gothic", 20, "bold"), text_color="#4CAF50")
right_top_title.pack(pady=(15, 8))

overview_text = CTkLabel(right_top_frame,
                         text="Sleep hours and trends will appear here.\n\n"
                              "Later, this section can show:\n"
                              "• Sleep duration\n"
                              "• Sleep quality\n"
                              "• Mood score",
                         font=("Century Gothic", 16), justify="center")
overview_text.pack(padx=30, pady=(5, 10))

right_bottom_frame = CTkFrame(contents_frame, fg_color=("gray95", "#212121"), corner_radius=20, width=520, height=220)
right_bottom_frame.place(relx=0.7, y=375, anchor="n")

right_bottom_title = CTkLabel(right_bottom_frame, text="Patterns & Insights",
                              font=("Century Gothic", 20, "bold"), text_color="#4CAF50")
right_bottom_title.pack(pady=(0, 8))

insight_text = CTkLabel(right_bottom_frame,
                        text="This area will compare wellbeing with progress.\n\n"
                             "Examples for your NEA:\n"
                             "• Mood vs sleep quality\n"
                             "• Sleep vs weekly weight trend",
                        font=("Century Gothic", 16), justify="center")
insight_text.pack(padx=30, pady=(5, 10))

def week_range_text(monday_date):
    start_str = monday_date.isoformat()
    end_str = (monday_date + timedelta(days=6)).isoformat()
    return f"Week: {start_str} to {end_str}"

moodnsleep.mainloop()
